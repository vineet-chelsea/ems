"""
Script to parse PM5320_modbus_used.xlsx and generate register mappings
"""
import sys
import json
import os

try:
    import openpyxl
except ImportError:
    print("Installing openpyxl...")
    os.system(f"{sys.executable} -m pip install openpyxl")
    import openpyxl

def parse_excel_file(file_path):
    """Parse Excel file and extract register mappings"""
    if not os.path.exists(file_path):
        print(f"Error: File not found: {file_path}")
        return None
    
    # Load with data_only=True to get calculated values instead of formulas
    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet = wb.active
    
    mappings = []
    
    # Find header row
    header_row = None
    for row_idx, row in enumerate(sheet.iter_rows(values_only=True), 1):
        row_values = [str(cell).lower() if cell else '' for cell in row]
        if any('parameter' in val or 'variable' in val for val in row_values):
            header_row = row_idx
            break
    
    if not header_row:
        print("Error: Could not find header row")
        print("First few rows:")
        for idx, row in enumerate(list(sheet.iter_rows(values_only=True))[:5], 1):
            print(f"Row {idx}: {[str(c) if c else '' for c in row]}")
        return None
    
    # Find column indices
    header_row_data = list(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))[0]
    headers = [str(cell).lower() if cell else '' for cell in header_row_data]
    param_col = None
    address_col = 2  # Column C (0-indexed = 2) - MODBUS ADDRESS column as per user requirement
    datatype_col = None
    desc_col = None
    
    for idx, header in enumerate(headers):
        if 'parameter' in header or 'variable' in header:
            param_col = idx
        # address_col is hardcoded to column C (index 2)
        elif 'type' in header or 'datatype' in header:
            datatype_col = idx
        elif 'description' in header or 'desc' in header:
            desc_col = idx
    
    if param_col is None:
        print("Error: Could not find Parameter column")
        print(f"Found columns: {headers}")
        return None
    
    print(f"Using column C (index {address_col}) for Modbus addresses")
    print(f"Parameter column: {param_col} (index)")
    
    # Parse data rows
    for row_idx, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
        param = row[param_col] if param_col < len(row) else None
        address = row[address_col] if address_col < len(row) else None
        
        if not param or not address:
            continue
        
        param = str(param).strip()
        address_str = str(address).strip()
        
        # Extract numeric address
        address_num = None
        if address_str.isdigit():
            address_num = int(address_str)
        elif '0x' in address_str.lower():
            address_num = int(address_str, 16)
        else:
            # Try to extract number from string
            import re
            match = re.search(r'\d+', address_str)
            if match:
                address_num = int(match.group())
        
        if not address_num:
            continue
        
        # Get data type
        data_type = 'FLOAT32'  # default
        if datatype_col is not None and datatype_col < len(row) and row[datatype_col]:
            data_type = str(row[datatype_col]).strip().upper()
            if data_type not in ['INT16U', 'INT32U', 'FLOAT32', 'UTF8', '4Q_FP_PF', 'INT16S']:
                # Infer from parameter name
                if 'PF' in param or 'powerfactor' in param.lower():
                    data_type = '4Q_FP_PF'
                elif 'energy' in param.lower():
                    data_type = 'INT32U'
                else:
                    data_type = 'FLOAT32'
        
        # Get description
        description = ''
        if desc_col is not None and desc_col < len(row) and row[desc_col]:
            description = str(row[desc_col]).strip()
        
        mappings.append({
            'parameter': param,
            'address': address_num,
            'dataType': data_type,
            'description': description
        })
    
    return mappings

if __name__ == '__main__':
    # Find Excel file
    excel_file = None
    possible_paths = [
        'PM5320_modbus_used.xlsx',
        '../PM5320_modbus_used.xlsx',
        '../../PM5320_modbus_used.xlsx',
    ]
    
    for path in possible_paths:
        if os.path.exists(path):
            excel_file = path
            break
    
    if not excel_file:
        print("Error: Could not find PM5320_modbus_used.xlsx")
        print(f"Searched in: {possible_paths}")
        sys.exit(1)
    
    print(f"Parsing {excel_file}...")
    mappings = parse_excel_file(excel_file)
    
    if mappings:
        print(f"Found {len(mappings)} parameter mappings")
        
        # Output as JSON
        output_file = 'pm5320_mappings.json'
        with open(output_file, 'w') as f:
            json.dump(mappings, f, indent=2)
        
        print(f"Saved to {output_file}")
        
        # Also print TypeScript format
        print("\nTypeScript format:")
        print("registerMappings: [")
        for mapping in mappings:
            print(f"  {{ parameter: '{mapping['parameter']}', address: {mapping['address']}, dataType: '{mapping['dataType']}', description: '{mapping['description']}' }},")
        print("]")
    else:
        print("No mappings found")
        sys.exit(1)

